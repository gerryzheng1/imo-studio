"""Read-only verification of the five UI-regression exports; no workbook rewrites."""
from pathlib import Path
import json,zipfile,xml.etree.ElementTree as E
import openpyxl

root=Path(__file__).parent
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
reports=[]
for path in sorted((root/'v131-generated').glob('*.xlsx')):
 ident=int(path.name.split('.')[0])
 original=next((root/'default-input'/'cleaned').glob(str(ident)+'.*xlsx'))
 with zipfile.ZipFile(path) as z,zipfile.ZipFile(original) as src:
  assert z.testzip() is None
  for part in z.namelist():
   if part.endswith(('.xml','.rels')):E.fromstring(z.read(part))
  wb=E.fromstring(z.read('xl/workbook.xml'))
  owb=E.fromstring(src.read('xl/workbook.xml'))
  sheets=wb.find('s:sheets',ns)
  assert [s.get('name')for s in sheets]==[s.get('name')for s in owb.find('s:sheets',ns)]
  visible=[s for s in sheets if s.get('state','visible')=='visible']
  assert len(visible)==(2 if ident==6 else 1)
  for name in src.namelist():
   if name.startswith('xl/media/'):assert z.read(name)==src.read(name)
  for index,sheet in enumerate(sheets):
   part=f'xl/worksheets/sheet{index+1}.xml'
   node=E.fromstring(z.read(part));old=E.fromstring(src.read(part))
   for kind in ['cols','pageMargins','mergeCells']:
    a=node.find('s:'+kind,ns);b=old.find('s:'+kind,ns)
    assert E.tostring(a)==E.tostring(b)if a is not None and b is not None else a is b
   setup=node.find('s:pageSetup',ns)
   assert setup.get('fitToWidth')=='1'and setup.get('fitToHeight')=='1'
   oldsetup=old.find('s:pageSetup',ns)
   if oldsetup is not None:
    for key,value in oldsetup.attrib.items():
     if key not in ['fitToWidth','fitToHeight','scale']:assert setup.get(key)==value
  styles=E.fromstring(z.read('xl/styles.xml'));oldstyles=E.fromstring(src.read('xl/styles.xml'))
  for kind in ['fonts','fills','borders','cellXfs']:
   a=styles.find('s:'+kind,ns);b=oldstyles.find('s:'+kind,ns)
   assert len(a)>=len(b)
   assert all(E.tostring(a[i])==E.tostring(child)for i,child in enumerate(b))
 book=openpyxl.load_workbook(path,read_only=True,data_only=False)
 assert book.sheetnames
 for sheet in book:
  list(sheet.iter_rows(values_only=True))
 if ident==4:
  sheet=book.worksheets[0]
  assert sheet['C11'].value=='Brisbane'
  assert sheet['F12'].value=='CARGO OPERATION,BUNKER,CREW CHANGE,SPARE PARTS'
  assert [sheet[f'B{i}'].value for i in range(11,21)]==list(range(1,11))
 reports.append({'file':path.name,'sheets':book.sheetnames,'passed':True})
 book.close()
assert len(reports)==5
(root/'v131-xlsx-verification.json').write_text(json.dumps({'passed':5,'reports':reports,'nativeOfficePrintVerified':False},ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS 5 XLSX: independent read-only open, XML/ZIP, original sheets/styles/images/merges, one-page setup, max-10 ports and full remarks')
