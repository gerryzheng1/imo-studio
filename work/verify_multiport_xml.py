from pathlib import Path
import json
import os
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

root = Path(__file__).parent
inputs = {p.name.split('.')[0]: p for p in (root / 'default-input' / 'cleaned').glob('*.xlsx')}
browsers = os.environ.get('IMO_BROWSERS', 'chrome,msedge').split(',')
files = sorted(file for file in (root / 'multiport-full-generated').rglob('*.xlsx') if file.relative_to(root / 'multiport-full-generated').parts[0] in browsers)
assert len(files) == len(browsers) * 30, len(files)
visible = 0
for file in files:
    original = openpyxl.load_workbook(inputs[file.name.split('.')[0]], read_only=True)
    book = openpyxl.load_workbook(file, read_only=True)
    assert book.sheetnames == original.sheetnames, file
    book.close()
    original.close()
    with zipfile.ZipFile(file) as archive:
        assert archive.testzip() is None, file
        xmls = {name: archive.read(name) for name in archive.namelist() if name.endswith('.xml') or name.endswith('.rels')}
        for name, data in xmls.items():
            ET.fromstring(data)
            assert '郑晓天'.encode() not in data and b'Zheng Xiaotian' not in data, (file, name)
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        workbook = ET.fromstring(xmls['xl/workbook.xml'])
        for sheet in workbook.findall('s:sheets/s:sheet', ns):
            if sheet.attrib.get('state', 'visible') != 'visible':
                continue
            # The preserved template sheet order corresponds to the original worksheet parts.
            index = list(workbook.find('s:sheets', ns)).index(sheet) + 1
            tree = ET.fromstring(xmls[f'xl/worksheets/sheet{index}.xml'])
            setup = tree.find('s:pageSetup', ns)
            assert setup is not None and setup.attrib.get('fitToWidth') == '1' and setup.attrib.get('fitToHeight') == '1', (file, sheet.attrib)
            visible += 1

report = {'passed': len(files), 'visibleSheets': visible,
          'checks': 'CRC and every XML, independent openpyxl read-only open, original sheet inventory, per-sheet one-page width/height, no application author credit in workbook',
          'files': [str(file.relative_to(root / 'multiport-full-generated')) for file in files]}
(root / 'multiport-full-xml-verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'PASS {len(files)} multi-port workbooks, {visible} visible original worksheets')
