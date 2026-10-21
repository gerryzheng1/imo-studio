import json,pathlib,zipfile,re,xml.etree.ElementTree as E,posixpath,datetime,sys
from openpyxl import load_workbook
P=pathlib.Path(__file__).parent
N={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
S=N['s'];R=N['r']; checks=[]
defaults='--defaults' in sys.argv
cfg=json.loads((P/('default-config.json' if defaults else 'config.json')).read_text(encoding='utf-8'))
template_root=P/'default-input'/'cleaned' if defaults else P/'templates'
generated_root=P/'v12-generated' if defaults else P/'generated'
def canon(e):
 if e is None:return None
 return (e.tag,tuple(sorted(e.attrib.items())),(e.text or '').strip(),tuple(canon(c) for c in e))
def read(z):
 wb=E.fromstring(z.read('xl/workbook.xml'));rels={r.get('Id'):r.get('Target') for r in E.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
 ss=[]
 if 'xl/sharedStrings.xml' in z.namelist():ss=[''.join(t.text or '' for t in si.findall('.//s:t',N)) for si in E.fromstring(z.read('xl/sharedStrings.xml'))]
 sheets=[]
 for s in wb.findall('s:sheets/s:sheet',N):
  target=rels[s.get('{'+R+'}id')];path=posixpath.normpath(posixpath.join('xl',target)) if not target.startswith('/') else target[1:];root=E.fromstring(z.read(path));cells={}
  for c in root.findall('s:sheetData/s:row/s:c',N):
   v=c.find('s:v',N);v=v.text if v is not None else ''
   if c.get('t')=='s':v=ss[int(v)]
   if c.get('t')=='inlineStr':v=''.join(x.text or '' for x in c.findall('.//s:t',N))
   cells[c.get('r')]=v or ''
  sheets.append({'name':s.get('name'),'path':path,'xml':root,'cells':cells})
 return wb,sheets,ss
def assert_equal(a,b,msg):
 assert a==b,msg
def address(ref):
 letters,row=re.match(r'([A-Z]+)(\d+)',ref).groups();col=0
 for ch in letters:col=26*col+ord(ch)-64
 return col,int(row)
def allowed(ref,sc):
 if ref in {f['cell'] for f in sc['fields']}|{sc.get('page'),sc.get('titlePage')}:return True
 c,r=address(ref)
 for region in sc.get('clear',[])+sc.get('table',{}).get('clear',[]):
  a,b=(region.split(':')+[region])[:2];c1,r1=address(a);c2,r2=address(b)
  if c1<=c<=c2 and r1<=r<=r2:return True
 return False
def emitted_records(sheets,sc,key):
 t=sc['table'];prefix=sc['name'];out=[]
 for sh in sheets:
  if sh['name']==prefix or sh['name'].startswith(prefix[:25]+' (') or re.sub(r' \(\d+\)$','',sh['name'])==prefix:
   for row in t['rows']:
    val=sh['cells'].get(t['columns'][key]+str(row),'')
    if val and val!='NIL':out.append(val)
 return out
for scenario in (['baseline','overflow','reduced'] if defaults else ['baseline','overflow','reduced','departure']):
 state=json.loads((P/(scenario+'-state.json')).read_text(encoding='utf-8'))
 for c in cfg['templates']:
  id=c['id'];p=next((generated_root/scenario).glob(str(id)+'.*.xlsx'))
  orig=template_root/c['filename']
  with zipfile.ZipFile(orig) as a,zipfile.ZipFile(p) as b:
   assert b.testzip() is None
   for part in b.namelist():
    if part.endswith('.xml') or part.endswith('.rels'):E.fromstring(b.read(part))
   wa,sa,ssa=read(a);wb,sb,ssb=read(b)
   assert_equal([s['name'] for s in sb[:len(sa)]],[s['name'] for s in sa],f'{scenario}/{id} original sheet order')
   for j,sh in enumerate(sb):
    original=sa[j] if j<len(sa) else next(s for s in sa if re.sub(r' \(\d+\)$','',sh['name'])==s['name'])
    for tag in ['cols','mergeCells','sheetPr','pageMargins','pageSetup','printOptions','headerFooter','rowBreaks','colBreaks']:
     assert_equal(canon(original['xml'].find('s:'+tag,N)),canon(sh['xml'].find('s:'+tag,N)),f'{scenario}/{id}/{sh["name"]}/{tag} changed')
    rowsa=[r.attrib for r in original['xml'].findall('s:sheetData/s:row',N)];rowsb=[r.attrib for r in sh['xml'].findall('s:sheetData/s:row',N)]
    assert_equal(rowsa,rowsb,f'{scenario}/{id}/{sh["name"]} row heights changed')
    sc=next(s for s in c['sheets'] if s['name']==original['name'])
    outcells={x.get('r'):x for x in sh['xml'].findall('s:sheetData/s:row/s:c',N)}
    for oldcell in original['xml'].findall('s:sheetData/s:row/s:c',N):
     if not allowed(oldcell.get('r'),sc):assert_equal(canon(oldcell),canon(outcells.get(oldcell.get('r'))),f'{scenario}/{id} unmapped cell changed '+oldcell.get('r'))
   oldnames=list(wa.find('s:definedNames',N) or []);newnames=list(wb.find('s:definedNames',N) or [])
   assert_equal([canon(x) for x in oldnames],[canon(x) for x in newnames[:len(oldnames)]],f'{scenario}/{id} original print areas / names')
   for dn in newnames:
    if dn.get('name') in ['_xlnm.Print_Area','_xlnm.Print_Titles']:
     index=int(dn.get('localSheetId'));assert sb[index]['name'].replace("'","''") in dn.text
   mutable={'xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml','xl/styles.xml','xl/sharedStrings.xml','xl/calcChain.xml',*[s['path'] for s in sa]}
   for rule in c.get('drawingRules',[]):
    if state['declarations'][rule['source']]=='records':
     mutable.add(rule['path']);dra=E.fromstring(a.read(rule['path']));drb=E.fromstring(b.read(rule['path']))
     filtered=[]
     for anchor in dra:
      ids=[x.get('id') for x in anchor.iter() if x.tag.endswith('}cNvPr')]
      if not any(x in rule['nilAnchorIds'] for x in ids):filtered.append(canon(anchor))
     assert_equal(filtered,[canon(x) for x in drb],f'{scenario}/{id} wrong NIL/line removal')
   for part in a.namelist():
    if part not in mutable:assert_equal(a.read(part),b.read(part),f'{scenario}/{id} unrelated part changed: '+part)
   # Original style indices remain intact; date styles may be appended.
   sta=E.fromstring(a.read('xl/styles.xml'));stb=E.fromstring(b.read('xl/styles.xml'))
   for tag in ['fonts','fills','borders','cellStyleXfs','cellStyles','dxfs']:
    assert_equal(canon(sta.find('s:'+tag,N)),canon(stb.find('s:'+tag,N)),f'{scenario}/{id} styles {tag}')
   xfa=sta.find('s:cellXfs',N);xfb=stb.find('s:cellXfs',N);assert_equal([canon(x) for x in xfa],[canon(x) for x in list(xfb)[:len(xfa)]],f'{scenario}/{id} original XF modified')
   for sh in sb:
    for cnode in sh['xml'].findall('s:sheetData/s:row/s:c',N):
     f=cnode.find('s:f',N)
     if f is not None:
      assert f.text and '#REF!' not in f.text
      assert cnode.find('s:v',N) is not None,'known formula missing new cache'
   calc=wb.find('s:calcPr',N);assert calc.get('fullCalcOnLoad')=='1' and calc.get('forceFullCalc')=='1'
   # No historical name/port/ship value remains, including unused shared strings.
   texts=[v for sh in sb for v in sh['cells'].values()]+ssb
   for old in ['XUE,FENG','Noumea','Shanghai','HOEGH TROTTER','H\u00d8EGH TROTTER','A90624647','CHLORPROMAZINE TAB 25MGCHLORPROMAZINE 10\'s']:
    assert not any(old in str(v) for v in texts),f'{scenario}/{id} old value leaked: '+old
   if id in [1,2]:
    assert not any(sb[1]['cells'].values()),f'{scenario}/{id} historical annex not empty'
    side='arrival' if id==1 else 'departure';t=c['sheets'][0]['table'];records=emitted_records(sb,c['sheets'][0],'name');expected=[]
    for r in state[side]:
     person=next(x for x in state['people'] if x['id']==r['personId']);expected.append(r['overrides'].get('name',person['name']))
    assert_equal(records,expected,f'{scenario}/{id} crew not complete')
    for sh in sb:
     if sh['name']=='Sheet1':continue
     assert_equal(sh['cells'].get('L31'),state['ship']['master'],f'{scenario}/{id} master cache')
    assert_equal(sb[0]['cells']['I5'],str((datetime.date.fromisoformat(state['voyage'][side+'Date'])-datetime.date(1899,12,30)).days),f'{scenario}/{id} date')
    first=state[side][0];person=next(x for x in state['people'] if x['id']==first['personId']);assert_equal(sb[0]['cells']['K9'],first['overrides'].get('passport',person['passport']),f'{scenario}/{id} passport text')
   if id==4:assert_equal(emitted_records(sb,c['sheets'][0],'port'),[x['port'] for x in state['ports']],f'{scenario}/ports')
   if id==5:
    vals=emitted_records(sb,c['sheets'][0],'name');assert_equal([v for v in vals if v.startswith('CURRENT STORE')],[x['name'] for x in state['stores']],f'{scenario}/stores')
    for x in state['stores']:assert x['location']+':' in vals
    assert not any(v.startswith('Storage: ') for v in vals)
   if id==6:
    vals=emitted_records(sb,c['sheets'][1],'illness');assert_equal(vals,[x['illness'] for x in state['cases']],f'{scenario}/cases')
   if id==7:assert_equal(emitted_records(sb,c['sheets'][0],'name'),[x['name'] for x in state['passengers']],f'{scenario}/passengers')
   if id==8:assert_equal(emitted_records(sb,c['sheets'][0],'name'),[x['name'] for x in state['medicines']],f'{scenario}/medicines')
   if id in [3,10]:
    side=state['voyage']['mode'];expected=[next(p for p in state['people'] if p['id']==r['personId'])['name'] for r in state[side]];assert_equal(emitted_records(sb,c['sheets'][0],'name'),expected,f'{scenario}/{id}/stable people')
   loaded=load_workbook(p,read_only=False,data_only=False,keep_links=True)
   assert_equal(loaded.sheetnames,[s['name'] for s in sb],f'{scenario}/{id} openpyxl sheet list')
   if id in [1,2]:
    assert loaded.worksheets[0]['K9'].data_type=='s'
    assert loaded.worksheets[0]['I9'].data_type=='s'
    assert loaded.worksheets[0]['K9'].number_format=='@'
   loaded.close()
   checks.append({'scenario':scenario,'id':id,'file':p.name,'sheets':len(sb),'preserved':True,'recordsComplete':True,'openpyxlRead':True})
(P/('v12-xml-verification.json' if defaults else 'xml-verification.json')).write_text(json.dumps({'files':checks,'count':len(checks)},ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS',len(checks),'workbooks: XML, original sheet order, merged regions, row heights, columns, print settings, fonts, borders, styles, untouched parts, formulas, current data, and no old records.')
